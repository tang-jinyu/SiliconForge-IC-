from pathlib import Path
from datetime import date
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist" / "documents" / "SiliconForge_数字IC可信设计智能体_比赛技术报告.docx"
ASSET = ROOT / "tmp" / "competition_report" / "assets"
ASSET.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

GREEN = "55B938"
DARK = "17221B"
PALE = "EAF5E6"
GRAY = "66736B"
LIGHT = "F5F7F5"

def font(size):
    paths = [Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf")]
    for p in paths:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()

def diagram(path, title, rows):
    w, h = 1800, 240 + len(rows) * 190
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.text((90, 60), title, font=font(52), fill=(23,34,27))
    box_w = (w - 180 - (len(rows[0])-1)*45)//len(rows[0])
    for ri, row in enumerate(rows):
        y = 180 + ri*190
        for ci, label in enumerate(row):
            x = 90 + ci*(box_w+45)
            d.rounded_rectangle((x,y,x+box_w,y+112), radius=20, fill=(234,245,230), outline=(85,185,56), width=4)
            bbox=d.textbbox((0,0),label,font=font(31)); tw=bbox[2]-bbox[0]
            d.text((x+(box_w-tw)/2,y+34),label,font=font(31),fill=(23,34,27))
            if ci < len(row)-1:
                ax=x+box_w+8; ay=y+56
                d.line((ax,ay,ax+28,ay),fill=(85,185,56),width=6)
                d.polygon([(ax+28,ay),(ax+16,ay-10),(ax+16,ay+10)],fill=(85,185,56))
        if ri < len(rows)-1:
            cx=w//2; y2=y+112
            d.line((cx,y2+10,cx,y2+63),fill=(85,185,56),width=6)
            d.polygon([(cx,y2+63),(cx-10,y2+47),(cx+10,y2+47)],fill=(85,185,56))
    im.save(path)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); trPr.append(e)

def set_keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def add_page_num(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run(); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); run._r.addnext(fld)

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"
    set_repeat_table_header(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,DARK); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True; r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: set_cell_shading(cells[i],LIGHT)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(8.7)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t

def add_pic(doc, path, caption, width=6.9):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(width))
    cp=doc.add_paragraph(caption); cp.style="Caption"; cp.alignment=WD_ALIGN_PARAGRAPH.CENTER

def add_bullets(doc, items):
    for s in items:
        p=doc.add_paragraph(style="List Bullet"); p.add_run(s)

def add_numbered(doc, items):
    for s in items:
        p=doc.add_paragraph(style="List Number"); p.add_run(s)

def add_manual_numbered(doc, items, compact=False):
    for i,s in enumerate(items,1):
        p=doc.add_paragraph()
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.left_indent=Pt(18)
        p.paragraph_format.hanging_indent=Pt(18)
        if compact:
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.space_after=Pt(1)
        p.add_run(f"{i}.  {s}")

def add_toc_grid(doc, items):
    half=(len(items)+1)//2
    t=doc.add_table(rows=half,cols=2)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    tblPr=t._tbl.tblPr
    borders=tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders=OxmlElement("w:tblBorders"); tblPr.append(borders)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"nil"); borders.append(e)
    for r in range(half):
        for c in range(2):
            idx=r+c*half
            cell=t.cell(r,c)
            if idx<len(items):
                p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                run=p.add_run(f"{idx+1}.  {items[idx]}"); run.font.size=Pt(8.5)
    return t

def add_section(doc, title, paras):
    doc.add_heading(title, level=1)
    for x in paras:
        if isinstance(x, tuple) and x[0]=="h2": doc.add_heading(x[1],level=2)
        elif isinstance(x, tuple) and x[0]=="bullets": add_bullets(doc,x[1])
        else: doc.add_paragraph(x)

architecture = ASSET / "architecture.png"
flow = ASSET / "flow.png"
govern = ASSET / "governance.png"
diagram(architecture,"SiliconForge 分层架构",[["中文 Web 工作台","HTTP API / 任务服务","本地队列与状态机"],["数字 IC 工作流编排","Prompt / 模板 / 反馈记忆","模型适配与超时降级"],["Icarus Verilog / Vivado","Yosys 综合","结构化产物与报告"]])
diagram(flow,"从赛题到可复核交付物的数据流",[["自然语言赛题","需求解析","接口契约冻结","RTL 与独立 TB"],["预检","仿真","综合","诊断与有界修复"],["质量门禁","中文报告","工作台审阅","整包下载 / 入库"]])
diagram(govern,"Verified Skills 对齐路线",[["Skill 目录与能力说明","静态扫描","隔离运行"],["独立评测集","技能卡 / 证据","签名与发布"]])

doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
normal=styles["Normal"]; normal.font.name="Microsoft YaHei"; normal.font.size=Pt(10.5); normal._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
normal.paragraph_format.line_spacing=1.35; normal.paragraph_format.space_after=Pt(5.5); normal.paragraph_format.first_line_indent=Pt(21)
for name,size,color in [("Title",30,DARK),("Heading 1",18,DARK),("Heading 2",13,GREEN),("Heading 3",11,DARK)]:
    st=styles[name]; st.font.name="Microsoft YaHei"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
    st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(7); st.paragraph_format.keep_with_next=True
styles["Caption"].font.name="Microsoft YaHei"; styles["Caption"].font.size=Pt(9); styles["Caption"].font.color.rgb=RGBColor.from_string(GRAY); styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"),"微软雅黑")
header=sec.header.paragraphs[0]; header.text="SiliconForge｜数字 IC 可信设计智能体比赛技术报告"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
add_page_num(sec.footer.paragraphs[0])

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(95); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("SiliconForge"); r.font.name="Aptos Display"; r.font.size=Pt(24); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(GREEN)
p=doc.add_paragraph(); p.style="Title"; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("数字 IC 可信设计智能体\n比赛技术报告")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16)
r=p.add_run("功能 · 架构 · 设计原理 · 工程验证 · 安全治理"); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(GRAY)
doc.add_paragraph("\n\n")
add_table(doc,["项目","内容"],[["报告版本","V1.0 / 2026-09-28"],["项目形态","可运行的数字 IC 自动化工作流与中文 Web 产品"],["目标平台","Windows 开发验证；Linux / 双 H100 迁移"],["技术定位","自然语言赛题到规格、RTL、TB、EDA、修复与报告的可信闭环"]],[1.5,5.0])
p=doc.add_paragraph("项目团队"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(36)
doc.add_page_break()

doc.add_heading("摘要",level=1)
doc.add_paragraph("SiliconForge 面向数字 IC 设计竞赛与中小型 RTL 开发，解决“大模型会写代码但难以形成可综合、可验证、可复现交付物”的核心矛盾。系统将自然语言需求转化为冻结的接口契约和设计规格，在同一任务上下文中生成 RTL 与自检 testbench，调用真实 EDA 工具完成预检、编译、仿真与综合，并把工具诊断送入受预算约束的修复循环。最终结果以 rtl、tb、doc、logs、reports 等目录组织，可在浏览器工作台中预览、编辑、比较、下载和沉淀为参数化模板。与普通对话式代码生成相比，本项目的交付对象不是一段“看起来合理”的 Verilog，而是一套具有可追溯证据链的工程包。")
doc.add_paragraph("工程已经完成中文 GUI、任务服务、同步与队列执行、OpenAI-compatible 模型适配、阶跃星辰调用、Icarus Verilog/Vivado 仿真、Yosys 综合、错误驱动修复、结构化报告、五类模板库、个人模板入库、盲测框架、超时重试与降级、Windows 到双 H100 的可移植部署骨架。实测方面，三项阶段一基线全部通过并获得平均 100 分；66×66 乘法器在 Vivado 仿真与 Yosys 综合中均通过；五类内置模板完成独立 EDA 回归。首轮五类端到端盲测为 1/5、平均 28 分，四项失败主要由模型服务 Connection error 引发并导致产物缺失。这一结果被保留为真实工程证据，说明系统已具备功能闭环，同时仍需通过本地 H100 模型、缓存与降级策略提升端到端可用性。")
p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.add_run("关键词：").bold=True; p.add_run("数字 IC；RTL 生成；EDA 闭环；大语言模型；参数化模板；Verified Skills；H100")

doc.add_heading("目录",level=1)
add_toc_grid(doc,["项目背景与比赛目标","需求与痛点分析","系统功能设计","总体架构与数据流","核心设计原理","Prompt、知识与反馈机制","模板库与个人能力资产","EDA 闭环与质量门禁","安全治理与 Verified Skills 对齐","实验设计与结果","Windows 产品化与双 H100 迁移","竞赛得分映射与演示方案","创新性、对比与论文可行性","关键模块实现说明","局限与后续路线","结论","参考文献","附录"])
doc.add_page_break()

add_section(doc,"1 项目背景与比赛目标",[
"生成式人工智能已经能够输出 Verilog 片段，但数字 IC 交付与一般软件代码生成存在本质差异：语法正确并不等于接口正确，仿真通过并不等于验证充分，可仿真也不等于可综合。复位极性、位宽符号、握手保持、FIFO 边界、状态机重叠匹配等细节，只要一个条件理解错误，就可能产生表面整洁却无法使用的实现。传统聊天界面还缺少文件组织、工具日志、修复历史和复现实验，难以满足比赛评审对技术深度、完整性和可展示性的要求。",
"本项目把“让模型写 RTL”重新定义为“让 Agent 对一个数字 IC 任务负责到可复核交付”。其目标不是替代工程师签核，而是压缩从需求澄清、规格固化、初版实现、测试平台、EDA 诊断到文档整理的重复劳动，使工程师把注意力集中在架构取舍、覆盖盲区和最终审阅。",
("h2","1.1 比赛导向"),
"根据比赛直播答疑与操作说明，评审同时关注实用性与创新性、技术深度、项目完整性、平台适配、演示效果和技术文章。系统因此采用“可运行产品 + 可解释架构 + 可量化实验 + 安全治理证据”的组合，而不是只提交模型调用脚本。界面提供一键闭环和可视化工作台，文档说明方法来源与边界，盲测报告展示成功和失败，部署方案说明 Windows 开发与 NVIDIA GPU 迁移之间的连续性。",
("h2","1.2 价值主张"),
("bullets",["对参赛者：输入赛题后快速得到可编辑、可验证、可下载的工程包。","对数字 IC 工程师：把常见模块与可靠修复经验沉淀为个人模板，降低重复劳动和 token 成本。","对评审：每个关键结论都能追溯到规格、代码、EDA 日志和报告，而非只依赖语言模型自述。","对平台生态：以技能封装、扫描、评测、签名为演进方向，使 Agent 能力可治理、可复用。"])
])

add_section(doc,"2 需求与痛点分析",[
("h2","2.1 从自然语言到硬件契约的鸿沟"),
"用户通常只会给出“设计一个 66×66 乘法器”“实现一个 FIFO”或一段竞赛描述，其中可能没有模块名、端口方向、复位语义、时序协议、溢出规则与允许延迟。直接生成 RTL 会把这些空白变成模型的隐式猜测。SiliconForge 在生成代码前先产出需求分析、架构草案和接口契约，把宽度、signed/unsigned、时钟与复位、握手、边界行为、可综合约束写成单一事实源。后续 RTL、TB、预检和修复均引用该契约，减少阶段之间的语义漂移。",
("h2","2.2 自证式 testbench 的局限"),
"若同一个模型同时误解了需求并生成 RTL 与 TB，两者可能在同一错误假设上保持一致，出现“错误地通过”。因此系统区分开发态自检 TB 与评测态独立 oracle。自检 TB 用于快速定位回归，隐藏 oracle 不进入 Agent prompt，只在生成结束后对 DUT 进行独立判定。该思想与 VerilogEval 以可执行测试衡量功能正确性的路线一致[1]。",
("h2","2.3 产物碎片与不可复现"),
"单纯复制聊天代码会产生文件名混乱、接口版本不一致、日志丢失和无法重跑等问题。系统采用固定输出目录、manifest、阶段 JSON、EDA 结果和中文说明书，使一次运行具备任务输入、配置、模型信息、代码、测试、诊断和结论。工作台只是在这些结构化事实之上的视图，避免页面状态成为唯一记录。",
("h2","2.4 模型服务并不稳定"),
"真实使用中，网络、配额、响应格式、超时和供应商兼容性都会造成失败。首轮盲测中 4 个任务因 Connection error 未产生完整产物，证明端到端可靠性必须被视为系统指标。项目增加统一适配层、超时、指数退避重试、失败分类、模板降级和本地 H100 路线；同时保留失败报告，防止用挑选成功样例替代可靠性评估。"
])

add_section(doc,"3 系统功能设计",[
("h2","3.1 一键完整闭环"),
"用户在中文网页输入赛题并点击提交，系统自动经历需求分析、架构设计、规格生成、RTL 生成、TB 生成、预检、仿真、综合、诊断、修复、审阅和报告。接口契约缺失时由 Agent 主动补全，不要求用户先填写专业表格；界面仍允许高级用户覆盖模块名、接口或输出目录。任务既可同步执行，也可提交本地队列，状态包含 queued、running、succeeded、failed、cancelled，并保留重试关系。",
("h2","3.2 专用任务入口"),
"除完整闭环外，系统支持模块生成、TB 修复、EDA 分诊和接口契约检查。模块生成适合从零创建小型 IP；TB 修复用于已有测试平台不稳定或刺激不足；EDA 分诊读取既有 RTL/TB 并给出编译、仿真、综合问题摘要；接口检查在进入昂贵模型调用前发现端口、位宽和命名不一致。多个入口复用同一任务服务、状态模型和产物层，避免形成相互割裂的脚本。",
("h2","3.3 类 VS Code 工作台"),
"工作台按 rtl、tb、doc、logs、reports 展示单一输出根目录。左侧为文件树，中部为带行号文本编辑器，右侧为任务与质量摘要；用户可查看源码、testbench、报告与日志，保存修改后重新验证，也可一键下载完整 ZIP。中文界面强调操作路径和结论，代码、工具命令及行业缩写保留原文。",
("h2","3.4 模板库与个人库"),
"系统内置组合运算、FSM、FIFO、协议控制器、修复任务五类参数化模板。命中明确模式时，可先实例化经过验证的模板，再执行独立 EDA，而非每次重新请求大模型。完成任务后，用户可以把高质量实现保存到个人库；入库前记录参数、接口、适用条件、验证证据和来源，使用时重新实例化并回归，防止把偶然成功的硬编码样例永久化。",
("h2","3.5 报告与下载"),
"每次任务交付不仅包含 Verilog 文件，还应包含需求摘要、接口表、设计原理、验证策略、EDA 结果、修复记录、已知限制和复现方式。这样比赛展示可以从功能演示自然切换到架构与证据，工程使用也能把结果直接交给同伴审阅。"
])

add_pic(doc,ROOT/"video_assets"/"competition"/"gui_03_template_library.png","图 1  中文工作台中的参数化模板库",6.9)

add_section(doc,"4 总体架构与数据流",[
"系统采用分层而非“一个超级 prompt”架构。展示层负责提交、进度、编辑和下载；任务服务统一校验输入、分配目录、持久化状态；工作流编排层控制阶段、预算和失败转移；能力层包含数字 IC 技能、Prompt、模板库和反馈记忆；模型适配层屏蔽 chat/completions 与 Responses 等传输差异；工具层执行真实 EDA；产物层保存可追溯文件。层间通过明确数据模型连接，使模型、EDA 或界面能够替换而不破坏核心流程。",
("h2","4.1 展示与服务层"),
"Web GUI 只通过 HTTP API 操作任务，不直接调用模型或 shell。请求模型包含 task_kind、execution、mode、dry_run、requirement_text、task_file、output_dir 和 metadata；服务端完成“需求文本或任务文件至少存在一个”等契约校验。任务记录保存创建时间、开始/结束时间、状态、错误、输出目录、重试来源和提供方请求标识，为取消、恢复、差异比较和审计提供基础。",
("h2","4.2 编排与能力层"),
"编排器把阶段输出写入上下文，但只传递下一阶段真正需要的信息，防止上下文无限膨胀。数字 IC Skill 定义可综合、自包含、接口冻结、独立验证、受限修复等规则；模板路由处理高频确定性任务；反馈记忆只沉淀经 EDA 证实的故障模式和修复约束，且按任务类别隔离。",
("h2","4.3 工具与产物层"),
"预检先做低成本结构检查，再调用 Icarus Verilog/Vivado 仿真和 Yosys 综合。工具命令由系统生成，需求文本和模型输出不被当作任意命令执行。每一轮工具返回退出码、标准输出、标准错误与摘要，编排器据此决定通过、修复或终止。产物层使用任务级目录避免相互覆盖，并以 manifest 建立文件类型和来源索引。"
])
add_pic(doc,architecture,"图 2  SiliconForge 分层架构",6.9)
add_pic(doc,flow,"图 3  从自然语言赛题到工程交付物的数据流",6.9)

add_section(doc,"5 核心设计原理",[
("h2","5.1 Contract First：先冻结契约，再生成代码"),
"接口契约是全链路的单一事实源，至少包含模块名、端口方向、宽度、时钟、复位、握手、输入输出时序、边界行为和不支持项。若需求有歧义，Agent 给出可审阅假设并记录在规格中；一旦冻结，RTL 与 TB 不得各自发明接口。预检可在模型调用后立即发现端口缺失、位宽不匹配、未定义模块和非自包含实现。",
("h2","5.2 生成与验证适度解耦"),
"RTL prompt 侧重可综合性、确定性赋值、状态转移、算术宽度与资源推断；TB prompt 侧重边界、时序、scoreboard、超时和自检。两者共享契约但采用不同检查清单。评测时再引入与模型上下文隔离的 oracle，使通过结果具有更强可信度。",
("h2","5.3 工具反馈优先于语言反馈"),
"AutoChip 的关键启示是把编译器和仿真器的真实输出返回模型，而不是让模型凭空反思；其论文报告在对应实验设置下可提升生成准确率[2]。本项目把语法错误、elaboration 错误、仿真超时、断言失败与综合错误规范化为诊断对象，只把相关片段和契约送入修复阶段。这样模型看到的是可操作证据，而不是笼统的“代码不对”。",
("h2","5.4 有界修复而非无限循环"),
"修复循环设置最大轮数、单轮超时和总任务预算。每轮保留原始文件、补丁、工具日志和结果，若同类错误连续出现则停止并给出人工建议。该策略避免 token 与算力无界消耗，也防止模型为了通过某个 TB 而破坏接口或删除检查。修复后必须重新运行完整质量门禁，不能仅凭模型声明成功。",
("h2","5.5 结构化交付优先"),
"目录、manifest 和报告不是附属品，而是 Agent 可用性的组成部分。任务结束条件定义为：必需文件存在、接口一致、仿真/综合达到目标、报告生成、失败也有清晰原因。即使模型调用失败，系统仍应生成任务记录和失败报告，以便复现和重试。"
])

add_section(doc,"6 Prompt、知识与反馈机制",[
("h2","6.1 分阶段 Prompt 体系"),
"系统不使用一段覆盖所有任务的长提示词，而是按照需求解析、规格、RTL、TB、修复、审阅分阶段组织。每个 prompt 包含角色、输入契约、硬约束、输出格式和失败处理；通用规则位于共享层，任务类别规则按组合逻辑、FSM、FIFO、协议和修复分开。提示词强调“只输出当前阶段需要的内容”，减少解释性文字污染代码。",
("h2","6.2 高质量 RTL 规则"),
"RTL 生成要求明确组合与时序边界，组合逻辑提供完整默认赋值，时序逻辑使用非阻塞赋值，参数计算避免零宽度和越界，signed 运算显式转换，状态机定义非法状态恢复，握手信号在 backpressure 下保持稳定，FIFO 同时处理读写与满空边界。禁止不可综合延时、initial 依赖、测试专用后门和未定义黑盒，除非规格明确允许。",
("h2","6.3 高质量 TB 规则"),
"testbench 必须自检而非仅产生波形：生成可重复激励，覆盖零值、全一、最大值、边界转换和随机样本；对时序协议建立 scoreboard；设置超时；失败打印期望值、实际值和上下文；成功输出唯一 TEST_PASS 标记。对于状态机使用重叠序列、复位中断和连续匹配；对于 FIFO 覆盖空读、满写、回卷和并发读写。",
("h2","6.4 文献与开源工程的转化"),
"VerilogEval 促使项目采用可执行基准与隐藏 oracle[1]；RTLLM 强调语法、功能和设计质量的多目标评价，并启发先规划后编码[3]；AutoChip 直接支持错误驱动修复[2]；ChipNeMo 表明领域适配可由 tokenizer、继续预训练、指令对齐和检索共同实现[4]，因此本项目把论文、Prompt 和模板作为检索资产，而非把所有知识硬塞进系统提示；ChatEDA 展示了 LLM 与 EDA executor 协同的可行性[5]，对应本项目的模型—工具分层；Reflexion 的语言反馈与情景记忆思想[6]被改造成“只有经过 EDA 证实的经验才允许入库”。",
("h2","6.5 反馈记忆的防污染设计"),
"反馈条目包含任务类别、触发症状、诊断证据、修复策略、适用边界和验证结果。仅有模型自评、单次成功或缺少工具日志的内容不能成为全局规则。读取时按模块类别和错误类型检索，避免 FIFO 的特殊修复被错误注入乘法器任务。未来可为每条经验增加版本、命中次数、成功率和撤销机制。"
])

add_section(doc,"7 模板库与个人能力资产",[
"模板库解决两个现实问题：常见数字电路没有必要每次从零生成；用户的高质量设计不应在一次会话后消失。内置五类模板覆盖频繁出现且容易参数化的基础任务，个人库则允许实验室或工程师围绕自身协议、编码规范和器件约束建立专属 Agent 资产。模板命中并不绕过验证，恰恰是把“生成不确定性”替换为“实例化 + 重新验证”。",
("h2","7.1 五类内置模板"),
])
doc.paragraphs[-1].paragraph_format.page_break_before=True
add_table(doc,["类别","典型参数","关键验证点","作用"],[
["组合运算 / ALU","WIDTH、操作集合、signed","进位、溢出、全边界","稳定处理位宽与组合逻辑"],
["FSM","状态编码、序列、重叠匹配","复位、非法态、连续命中","约束状态转移与输出时序"],
["同步 FIFO","DATA_WIDTH、DEPTH","满空、指针回卷、并发读写","复用可靠存储控制骨架"],
["协议控制器","数据宽度、采样沿、帧格式","握手保持、位计数、帧边界","复用协议时序模式"],
["修复任务","故障类型、目标约束","修复前后差异、回归","把常见错误变为可复核补丁"]
],[1.1,1.5,2.2,2.0])
add_section(doc,"7.2 个人模板入库流程",[
"用户在工作台选择一个已通过的运行结果，系统提取 RTL/TB、接口契约和参数候选，要求填写模板名称、用途和适用边界。入库前重新执行独立 EDA；保存时记录来源任务、工具版本、通过时间和内容哈希。调用时只替换声明的参数，不允许字符串任意拼接改变结构；实例化结果进入新任务目录并再次仿真与综合。未来还应增加模板版本、废弃标记、团队审核和签名。",
"模板路线的直接收益是降低延迟、token 消耗和供应商故障影响；更深层价值是让 Agent 从公共大模型变为具有个人工作场景记忆的工程助手。其创新点不在于保存代码片段，而在于把代码、参数、契约、验证与来源绑定成可治理资产。"
])

add_section(doc,"8 EDA 闭环与质量门禁",[
("h2","8.1 分级门禁"),
"质量门禁由低成本到高成本逐级执行。第一层检查文件存在、模块唯一性、端口、未定义模块、明显测试注入和不可综合结构；第二层编译与 elaboration；第三层运行自检 TB；第四层综合并检查 latch、multiple driver 或工具错误；第五层生成审阅摘要。任何层失败都产生结构化诊断，只有全部满足任务目标才标记 overall_pass。",
("h2","8.2 工具链"),
"Windows 开发环境可使用 Icarus Verilog/vvp 进行快速仿真，以 Yosys 或 yowasp-yosys 进行开源综合；项目也已经在 66×66 乘法器案例中使用 Vivado 仿真日志作为附加证据。服务器迁移后保持同一抽象接口，可接入容器内工具，后续扩展 Verilator、Surelog、lint、formal、coverage、STA 与厂商综合。工具后端、版本和命令必须写入报告，避免不同机器上的“通过”含义不一致。",
("h2","8.3 输出目录"),
])
add_table(doc,["目录/文件","内容","审阅价值"],[["rtl/","DUT 与必要 helper 模块","查看可综合实现"],["tb/","自检 testbench、可选 oracle","审查刺激与判据"],["doc/","规格、接口表、设计说明、最终报告","理解设计意图与边界"],["logs/","模型与 EDA 日志","定位失败和复现"],["reports/","仿真、综合、盲测与质量摘要","快速判断通过状态"],["manifest.json","文件索引、来源、哈希与元数据","追溯完整性"]],[1.4,3.0,2.3])
add_section(doc,"8.4 修复策略",[
"修复 prompt 只接收冻结契约、当前 RTL/TB 的必要片段和规范化错误，要求输出最小修改。若错误来自 TB，禁止通过修改 DUT 迎合错误期望；若错误来自接口契约，必须回到需求阶段而非在 RTL 中隐式改口。修复完成后重新走预检、编译、仿真和综合，并在 repair_history 中记录轮次、根因、补丁和结果。"
])

add_section(doc,"9 安全治理与 Verified Skills 对齐",[
"NVIDIA Skills 文档把技能视为可移植的指令与能力封装，Verified Skills 的价值不是给 prompt 加一个“安全”标签，而是让能力经过目录化、扫描、独立评测、签名和技能卡说明，从而把主观声明转化为可验证证据[7–9]。对于能执行 EDA、写文件和调用外部模型的 Agent，这种治理尤其重要：错误或恶意技能可能读取密钥、越权写目录、执行需求中的命令，或者以看似合理的报告掩盖失败。",
("h2","9.1 三层安全模型"),
"模型层控制模型来源、系统指令和供应商边界；运行层限制文件系统、网络、进程和资源；能力层治理 Skill、模板、工具白名单和版本。三层不可互相替代：安全模型不能阻止一个被允许的危险 shell 命令，沙箱也不能保证技能在功能上有效，签名只能证明来源与完整性而不能证明质量。",
("h2","9.2 当前项目的对齐情况"),
])
add_table(doc,["治理要素","当前实现","状态","下一步"],[
["目录化","独立 digital-ic Skill、规则和版本化工程","已具备","生成正式技能卡"],
["扫描","任务目录约束、凭据排除、命令白名单与静态检查","部分具备","接入官方 SkillSpector/等价扫描"],
["评测","五类盲测、隐藏 oracle、EDA 报告","已具备基础","扩充规模并做有/无 Skill 对照"],
["签名","尚未生成官方 detached OMS signature","未完成","服务器发布阶段签名与验签"],
["文档","功能边界、用户说明、迁移说明、本文档","已具备","自动绑定哈希与版本"]
],[1.1,2.7,1.0,2.0])
add_section(doc,"9.3 采用该架构的实际区别",[
"不用治理框架时，任何人都可以复制一个 prompt 并声称“适用于数字 IC”，用户无法知道它读写了什么、是否跑过测试、何时更新、是否被篡改；采用后，每个 Skill 有明确权限、输入输出、评测基线和版本证据，更新必须重新扫描和评测。对本项目而言，这会把参赛作品从“能演示的 Agent”推进到“可在实验室共享和迭代的工程能力包”。当前报告只声明对齐路线，不宣称已经获得 NVIDIA 官方 Verified 标识。"
])
add_pic(doc,govern,"图 4  Verified Skills 治理路径及本项目演进方向",6.9)

add_section(doc,"10 实验设计与结果",[
("h2","10.1 实验原则"),
"实验区分确定性基线、真实模型端到端任务、模板回归和隐藏 oracle 盲测。指标至少包括任务成功、仿真、综合、overall_pass、修复轮数、必需产物和独立 oracle；后续还应加入 Pass@1、修复后通过率、平均时延、token、失败类型、PPA 与跨模型方差。报告保留失败，不通过手工改代码后只展示成功结果。",
("h2","10.2 当前可复现实测"),
])
add_table(doc,["实验","规模/对象","结果","证据解释"],[
["阶段一基线","加法、握手 FSM、同步 FIFO 共 3 项","3/3 通过，平均 100 分","仿真与综合均通过"],
["66×66 乘法器","mult66x66 端到端运行","overall_pass=true","Vivado 仿真 TEST_PASS，Yosys 综合通过"],
["五类内置模板","组合、FSM、FIFO、协议、修复","本轮独立 EDA 回归均通过","说明常见任务可零 token 实例化后验证"],
["五类首轮盲测","ALU、FSM、FIFO、SPI、修复","1/5，通过率 20%，平均 28 分","ALU 100；其余多因 Connection error 与产物缺失"],
["自动化测试","模板、提交契约、产物工作台","8 项测试通过","覆盖 API/目录/模板关键路径"]
],[1.35,1.5,1.45,2.45])
add_section(doc,"10.3 结果解读",[
"三项基线、66×66 乘法器与五类模板回归说明工程主链路、EDA 集成和结构化交付已经可用。尤其 66×66 任务验证了用户只提供高层需求时，系统可以补全接口与测试并形成真实工具证据，而不是依赖手工填写接口契约。模板回归说明常见设计可通过参数化资产稳定复用。",
"首轮盲测的低通过率不能被简单解释为 RTL 能力不足：4 个失败记录的主要症状是模型服务 Connection error，随后缺少必需产物。它揭示的是产品级端到端可靠性短板——网络和模型服务也是系统的一部分。下一轮应在相同题集上比较阶跃星辰与双 H100 本地模型，分别统计模型错误、传输错误、验证错误和修复结果；同时把模板降级纳入故障注入测试。",
"现有五个盲测案例用于建立类别覆盖和流水线回归，不能证明“这五类任意题目都能准确生成”。要形成论文级结论，需要每类扩展到至少数十个不同接口、位宽和边界条件的保留测试集，固定随机种子，报告置信区间，并设置无 Skill、仅 prompt、完整闭环、模板增强等消融组。"
])

add_pic(doc,ROOT/"video_assets"/"competition"/"gui_05_quality.png","图 5  工作台中的质量与验证视图",6.9)

add_section(doc,"11 Windows 产品化与双 H100 迁移",[
("h2","11.1 Windows 先行的合理性"),
"当前阶段优先在 Windows 把交互、任务契约、EDA、目录、报告和回归测试跑通，可以缩短产品迭代反馈。用户电脑没有适合部署大规模本地模型的 GPU 并不妨碍开发：模型可通过阶跃星辰等兼容 API 远程调用，RTL 仿真与小规模综合由本机 CPU/工具完成。密钥只存环境变量，不写入源码、文档和迁移包。",
("h2","11.2 可移植设计"),
"项目代码不依赖固定盘符；模型、base URL、transport、超时、温度和访问口令均环境化；历史 runs、虚拟环境、个人密钥和 Windows EDA 套件从压缩包排除。Docker Compose 负责 Linux 服务，统一 HTTP API 和 OpenAI-compatible 模型接口使前端与工作流无需因模型迁移而改写。Windows 打包脚本生成 ZIP 与 SHA256，便于把同一版本发送到服务器。",
("h2","11.3 双 H100 方案"),
"服务器安装 NVIDIA 驱动、Docker 与 NVIDIA Container Toolkit 后，以 vLLM 或同类推理服务暴露 OpenAI-compatible API，使用 tensor parallel size=2 分布到两张 H100。模型应实际承担 RTL/TB 生成、诊断或修复，并在运行报告中记录模型名、GPU、并行参数、时延和盲测结果，才能形成有意义的 NVIDIA 平台适配证据。若选择 30B–70B 级代码模型，应根据显存、上下文和并发设置精度、KV cache 与最大序列长度，而不是只追求参数量。",
("h2","11.4 迁移验收"),
("bullets",["/api/runtime 能识别 Linux、两张 H100、Icarus Verilog 与 Yosys。","阶段一确定性基线保持 3/3 通过。","同一盲测集分别运行远程 StepFun 与本地模型，保存可比报告。","验证断网时模板与本地模型降级，验证超时、重试、取消和恢复。","确认容器只挂载任务目录，密钥不进入镜像与日志，发布包哈希可复核。"])
])

add_section(doc,"12 竞赛得分映射与演示方案",[
"作品展示采用“痛点—一键功能—架构原理—研究依据—量化结果—平台迁移”的叙事。痛点回答为什么需要产品；完整屏幕录制证明 GUI 可操作；架构与闭环体现技术深度；参考文献说明设计不是任意堆叠；成功和失败数据体现严谨；双 H100 路线与 Verified Skills 对齐回应 NVIDIA 平台价值。",
])
add_table(doc,["评分维度","项目证据","演示重点"],[
["实用性与创新性","自然语言到工程包、模板个人库、零 token 路由","66×66 一键生成与下载"],
["技术深度","契约优先、独立验证、真实 EDA、有界修复","展示日志、repair 与架构图"],
["完整性","中文 GUI、任务服务、目录、报告、打包迁移","从输入到工作台闭环"],
["平台适配","Verified Skills 对齐、Docker、双 H100 本地模型路线","展示 runtime 与模型对照实验"],
["演示效果","全屏无变形录制、真实交互、关键数据字幕","不使用伪造进度与挑选截图"],
["技术文章","本报告、开发日志、引用与复现步骤","提供文档与代码链接"]
],[1.25,3.15,2.35])
add_section(doc,"12.1 推荐现场流程",[
"第一分钟用一个失败案例说明普通聊天生成的痛点；第二至四分钟在 GUI 输入“设计一个无符号 66×66 乘法器”，展示 Agent 自动补全接口、执行阶段进度和真实 EDA；第五分钟打开工作台，查看 rtl、tb、doc 与日志，编辑一行并重新验证；第六分钟打开模板库，说明五类模板和个人入库；第七分钟展示基线、盲测和错误分类；最后用架构图、Verified Skills 映射和双 H100 对照实验收束。现场必须准备离线模板演示和预生成结果，以防外部模型连接异常，同时明确标注哪些是实时运行、哪些是历史证据。"
])

doc.add_page_break()
add_section(doc,"13 创新性、对比与论文可行性",[
("h2","13.1 与已有路线的差异"),
])
add_table(doc,["路线","优势","不足","SiliconForge 的吸收与扩展"],[
["通用对话助手","交互自然、覆盖面广","无固定契约、EDA 与产物治理","保留自然语言入口，增加工程闭环"],
["VerilogEval/RTLLM","评测可重复、关注 RTL 质量","主要是基准或生成研究","引入生产式任务状态与交付工作台"],
["AutoChip","工具反馈修复有效","个人资产与 GUI 不是重点","加入有界修复、模板库和报告"],
["ChatEDA","LLM 与 EDA executor 协同","更偏流程自动化研究","扩展到需求—RTL—验证—文档闭环"],
["Verified Skills","能力治理与可信发布","不直接解决 RTL 生成质量","把数字 IC Skill、评测与签名路线结合"]
],[1.15,1.65,1.85,2.1])
add_section(doc,"13.2 可投稿研究问题",[
"本工程具有论文基础，但目前更适合把系统作为实验平台，而不是直接以“做了一个 GUI”投稿。最有潜力的方向是“检索增强的参数化 RTL Agent：模板复用、反馈记忆与 EDA 闭环如何共同提升正确率、稳定性和成本效率”。核心假设可以是：契约优先减少接口错误，独立 oracle 降低自证偏差，经验检索提升修复成功率，参数化模板显著降低延迟与 token，同时保持跨参数正确性。",
"实验需扩充为分层数据集：五类各 30–100 题，训练/开发/保留测试严格隔离；比较裸模型、通用 prompt、分阶段 prompt、+EDA 修复、+反馈记忆、+模板路由；至少使用两种模型和多次随机运行；报告 Pass@1、Pass@k、修复增益、token、时延、模板命中率、错误类型与可综合质量。对模板任务还要测试未见参数和结构变体，防止仅记住样例。",
"第二个可发表方向是“面向硬件生成 Agent 的可信技能治理”，研究 Skill 扫描、隐藏评测、来源签名和任务级沙箱如何降低提示注入、越权工具调用和能力回归。该方向需要构建攻击与误用数据集，并做安全性—可用性权衡。若只描述 NVIDIA 框架而无新机制，创新不足；若能提出硬件领域的权限模型、证据格式与持续验证方法，则具有研究价值。"
])

add_section(doc,"14 关键模块实现说明",[
("h2","14.1 需求解析与规格冻结"),
"需求解析器首先判断任务属于全新设计、局部修复还是既有工程分诊，并抽取功能、接口、时序、参数、复位、异常和交付约束。对于“设计一个 66×66 乘法器”这类简短输入，系统会推导两个 66 位无符号输入、132 位输出和纯组合实现作为默认方案，同时把“无流水、无符号、无时钟”写入假设列表；如果用户明确要求流水级数、吞吐或 signed，则覆盖默认值。所有推导都进入结构化规格，后续阶段只能引用而不能静默改变。这样既保留一键体验，又让隐含假设可被工程师检查。",
"规格冻结并不意味着拒绝修正。若编译或用户审阅发现契约本身错误，系统会新建规格版本并使原 RTL/TB 失效，随后重新生成相关阶段；若只是实现错误，则保持契约版本不变并执行局部修复。版本关系让评审能够区分“需求发生改变”和“代码没有遵守需求”，也防止 Agent 在失败后通过改写验收条件自我证明。",
("h2","14.2 架构设计与任务规划"),
"架构阶段在写代码前决定组合或时序实现、数据通路与控制通路边界、存储结构、流水策略和参数推导。对简单组合模块，规划器遵循 keep simple：不引入无意义状态机和寄存器；对 FIFO，明确读写指针、计数器或额外相位位方案；对协议控制器，绘制帧状态与采样事件；对修复任务，先复现错误并冻结正确接口。RTLLM 的 self-planning 思路说明，先组织设计步骤有助于避免模型在长代码中失去整体约束[3]，本项目进一步要求规划产物能够被后续规则与报告直接消费。",
"规划器还负责选择执行路径。若任务与受信模板的参数空间完全匹配，进入模板实例化；若匹配类别但结构有新要求，使用模板作为检索示例但仍由模型生成；若属于未知复杂任务，则进入通用多阶段工作流。路由结果、置信依据和降级原因写入运行记录，使低 token 的快速路径不会成为不可解释的黑箱。",
("h2","14.3 五类任务的领域策略"),
"组合运算类重点处理位宽传播、截断、进位、符号扩展和默认分支。生成前按表达式计算中间结果宽度，避免 Verilog 上下文宽度导致高位丢失；测试同时覆盖零、最大值、交替位型和随机向量。对乘法器，结果宽度默认取两个操作数宽度之和，除非规格明确要求截断或饱和。",
"FSM 类把状态、输入事件、输出语义和转移条件分表表示。Moore/Mealy 选择必须与输出时序一致，序列检测明确是否允许重叠，复位态和非法态回退必须存在。TB 不只走主路径，还注入中途复位、连续匹配、错误前缀和边界周期，防止“能看到一次正确波形”被误判为状态机正确。",
"FIFO 类的难点不是存储数组本身，而是满空判定、读写同周期语义和非 2 次幂深度。模板对 DEPTH 做参数约束，指针宽度与回卷逻辑显式计算；规格明确空读、满写是忽略、报错还是保持输出；scoreboard 维护软件队列并在每个有效读周期比较。若未来扩展异步 FIFO，还需加入跨时钟 Gray 指针、同步器和 CDC 验证，不能复用同步 FIFO 的通过结论。",
"协议控制器类把外部信号沿、帧边界、位序、握手保持和超时作为一级规格。例如 SPI 接收器必须明确 CPOL/CPHA、MSB/LSB first、片选行为和数据有效脉冲；ready/valid 控制器在 valid 已拉高而 ready 未到时必须保持 payload。模板仅覆盖声明的协议子集，未覆盖模式必须走模型生成和完整回归。",
"修复任务类坚持先复现、再诊断、后修改。系统比较失败日志与冻结契约，判断根因属于 RTL、TB、接口还是环境；优先修改最小范围并保留 diff。典型可复用规则包括补全组合默认赋值、修复非阻塞赋值时序、纠正位宽截断、延长不足的仿真预算和移除伪 direct injection。任何修复都不能通过删除断言、缩小测试或硬编码测试向量来获得通过。",
("h2","14.4 模型适配、重试与降级"),
"模型适配层把消息、温度、输出上限、超时和响应解析统一为内部请求。transport=chat 表示使用 chat/completions 风格的消息接口；在提供方支持时也可使用 Responses 等传输。选择 transport 只影响调用协议，不改变上层数字 IC 任务。系统对连接超时、读超时、限流、服务端错误和内容解析失败分别分类；只有幂等请求可自动重试，并使用带抖动的指数退避，避免瞬时故障造成并发雪崩。",
"当重试耗尽时，降级顺序遵循可解释原则：精确命中受信模板则实例化并验证；已配置本地 H100 模型则切换兼容端点；否则生成失败报告并保留重试入口。系统不会在未知任务上悄悄返回与需求不符的占位代码。provider、尝试次数、各轮耗时和最终路径都会进入 metadata，使盲测可以把网络可靠性与 RTL 能力分开分析。",
("h2","14.5 产物报告与证据链"),
"最终中文报告由结构化阶段事实生成，而不是让大模型自由编写成功故事。接口表来自冻结契约，文件清单来自 manifest，仿真与综合状态来自工具退出码和解析结果，修复记录来自 repair_history，限制来自规划与审阅。模型可以帮助解释原因和形成可读摘要，但不能覆盖工具事实。若 overall_pass 为 false，报告标题和结论必须明确失败，并列出缺失文件与建议动作。",
"每个产物都应具有角色和来源：需求文件说明输入，spec 说明承诺，RTL 是实现，TB 是开发态验证，oracle 是独立评测，日志是原始证据，报告是视图，manifest 是索引。下载 ZIP 保留相对路径与元数据；在模板入库或服务器迁移时，可用哈希检查内容是否改变。后续接入签名后，技能包、模板包和竞赛发布包都可形成来源—版本—验证—签名的连续链。",
("h2","14.6 可观测性与异常处理"),
"工程任务常常持续数分钟，用户需要知道系统究竟在分析、调用模型、运行仿真还是等待重试。服务层因此记录阶段开始、结束、进度、最近日志和取消标志；GUI 轮询或订阅状态并显示可理解的中文信息。取消操作在本地停止后续阶段，在提供方支持 request id cancel 时同时请求取消远端生成；已经写出的文件保留并标记未完成，便于问题定位。",
"错误信息面向两类读者：普通用户看到“模型连接失败，已重试三次，可使用模板或稍后重试”；工程师可展开查看异常类型、HTTP 状态、工具命令、退出码和日志路径。密钥、Authorization 头与敏感环境变量在日志中脱敏。异常处理的目标不是消灭失败，而是让失败有边界、可解释、可恢复。"
])

add_section(doc,"15 局限与后续路线",[
"第一，当前盲测规模小，不能外推到复杂 SoC、时序收敛或签核级验证；第二，评测工具主要集中在仿真和逻辑综合，尚未覆盖 lint、CDC、formal、coverage、STA 与 LEC；第三，模型服务连接仍是主要可靠性风险；第四，个人模板的自动参数抽取与安全审核尚需加强；第五，官方 Verified Skills 扫描、评测器和 OMS 签名未完成；第六，双 H100 只完成部署方案，尚未用本地模型跑出对照数据。",
"近期优先级应为：完成双 H100 推理服务和盲测对照；将五类题库扩展为分层保留集；把 provider 错误、模型错误、EDA 错误分开统计；接入 lint 与 formal 的最小子集；给模板增加版本、哈希、适用边界和回退；生成技能卡并接入扫描。中期再扩展 ASIC 任务、IP 集成、多 worker 调度、团队模板仓库和基于覆盖率的验证增强。"
])

add_section(doc,"16 结论",[
"SiliconForge 已经从早期的 Prompt 与脚本演进为可在浏览器操作的数字 IC 可信设计智能体。其核心贡献不是单次生成某个 Verilog，而是建立从赛题理解、接口冻结、RTL/TB 生成、真实 EDA、错误驱动修复到结构化报告的连续证据链；再通过参数化模板和个人库，把高频设计转化为可复用、低成本的工程资产；通过 Verified Skills 对齐和双 H100 迁移路线，为能力治理与本地部署预留可执行路径。",
"现有结果证明主链路可以交付，但首轮盲测也诚实暴露了外部模型连接与端到端可靠性的不足。比赛阶段的关键不应是继续堆砌功能名，而是把 H100 对照实验、盲测扩容、安全证据和现场演示做实。只要后续数据能够证明闭环、模板和反馈机制在正确率、成本与稳定性上的增益，本项目不仅具备线下展示价值，也具备围绕 RTL Agent 评测与可信工程化开展论文研究的可行性。"
])

doc.add_heading("参考文献",level=1)
refs=[
"[1] Liu, M., Pinckney, N., Khailany, B. & Ren, H. VerilogEval: Evaluating Large Language Models for Verilog Code Generation. ICCAD (2023). arXiv:2309.07544. https://arxiv.org/abs/2309.07544",
"[2] Thakur, S. et al. AutoChip: Automating HDL Generation Using LLM Feedback. arXiv:2311.04887 (2023). https://arxiv.org/abs/2311.04887",
"[3] Lu, Y., Liu, S., Zhang, Q. & Xie, Z. RTLLM: An Open-Source Benchmark for Design RTL Generation with Large Language Model. ASP-DAC (2024). arXiv:2308.05345. https://arxiv.org/abs/2308.05345",
"[4] Liu, M. et al. ChipNeMo: Domain-Adapted LLMs for Chip Design. arXiv:2311.00176 (2023). https://arxiv.org/abs/2311.00176",
"[5] He, Z. et al. ChatEDA: A Large Language Model Powered Autonomous Agent for EDA. IEEE Transactions on Computer-Aided Design of Integrated Circuits and Systems (2024). doi:10.1109/TCAD.2024.3383347; arXiv:2308.10204.",
"[6] Shinn, N. et al. Reflexion: Language Agents with Verbal Reinforcement Learning. arXiv:2303.11366 (2023). https://arxiv.org/abs/2303.11366",
"[7] NVIDIA. NVIDIA Skills Documentation. https://docs.nvidia.com/skills （访问日期：2026-09-28）.",
"[8] NVIDIA. Agent Skill Trust Pipeline. https://docs.nvidia.com/skills/agent-skill-trust-pipeline （访问日期：2026-09-28）.",
"[9] NVIDIA. Scanning Agent Skills. https://docs.nvidia.com/skills/scanning-agent-skills （访问日期：2026-09-28）.",
"[10] NVIDIA. Secure Agent Workspace Reference Design: Security and Governance Model. https://docs.nvidia.com/enterprise-reference-architectures/secure-agent-workspace-reference-design/latest/security-and-governance-model.html （访问日期：2026-09-28）.",
"[11] Wolf, C. et al. Yosys—A Free Verilog Synthesis Suite. Austrian Workshop on Microelectronics (2013). https://yosyshq.net/yosys/",
"[12] Icarus Verilog Project. Icarus Verilog Documentation. https://steveicarus.github.io/iverilog/ （访问日期：2026-09-28）.",
]
for ref in refs:
    p=doc.add_paragraph(ref); p.paragraph_format.first_line_indent=Pt(-18); p.paragraph_format.left_indent=Pt(18); p.paragraph_format.space_after=Pt(6)

doc.add_heading("附录 A 复现与验收清单",level=1)
add_manual_numbered(doc,["在工程根目录创建虚拟环境并安装项目依赖。","配置模型环境变量；API Key 只保存在本机环境文件。","启动服务后打开 http://127.0.0.1:8000/，先运行阶段一演示。","输入 66×66 乘法器需求，确认任务目录包含 rtl、tb、doc、logs、reports。","检查 eda_result.json 的 simulation_passed、synthesis_passed 和 overall_pass。","运行五类盲测并保存 benchmark_report.json；失败不得删除。","执行自动化测试并记录工具与模型版本。","迁移服务器后按同一清单复测，并生成 StepFun 与 H100 本地模型对照表。"])
doc.add_heading("附录 B 术语与边界",level=1)
add_table(doc,["术语","本文含义"],[["闭环","生成结果进入真实工具，诊断可触发受限修复，最后形成报告"],["通过","满足该任务配置的仿真、综合与必需产物门禁；不等于芯片签核"],["隐藏 oracle","不进入 Agent prompt 的独立测试平台或判据"],["零 token 路由","模板实例化阶段不调用大模型，但仍执行 EDA"],["Verified 对齐","遵循目录化、扫描、评测、签名、技能卡路线；不等于已获官方认证"],["H100 支持","具备容器和兼容 API 迁移方案；实测对照尚待服务器完成"]],[1.4,5.2])

# prevent table rows splitting and ensure CJK font everywhere
for table in doc.tables:
    for row in table.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); trPr.append(cant)
for p in doc.paragraphs:
    for r in p.runs:
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"微软雅黑")

doc.core_properties.title="SiliconForge 数字 IC 可信设计智能体比赛技术报告"
doc.core_properties.subject="功能、架构、设计原理、工程验证与安全治理"
doc.core_properties.author="SiliconForge 项目团队"
doc.core_properties.keywords="数字IC, RTL, Agent, EDA, Verified Skills, H100"
doc.save(OUT)

text="\n".join(p.text for p in doc.paragraphs)
cn=len(re.findall(r"[\u4e00-\u9fff]",text))
print(f"written={OUT}")
print(f"chinese_chars={cn}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
